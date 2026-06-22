import sys
import re
import json
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
RESULTS_DIR = BASE / "results_cil"
DIST_XLSX = BASE / "data/partitions/100_client/label_skew_0.1/label_skew_0.1_distribution.xlsx"
TASK_FILE = BASE / "results/incremental_order/label_skew_0.1_10_client_5.0_task.txt"
LABEL_FILE = BASE / "data/label_mapping.txt"


#"FedAvg_finetune_100client_label_skew_0.1_5.0_gru.log",
#"FedAvg_lwf_100client_label_skew_0.1_5.0_gru_alpha0.99_temp2.0.log",
#"FedAvg_ewc_100client_label_skew_0.1_5.0_gru_lambda1000.0.log",
#"FedAvg_ours2_100client_label_skew_0.1_5.0_gru_mem1.0_kg1.0_eva0.95_blom0.9_archive.log"
#"FedAvg_mas_100client_label_skew_0.1_5.0_gru_lambda200.0.log",
#"FedAvg_feat_100client_label_skew_0.1_5.0_gru_mem1.0_lam0.1_rho0.9_temp0.5.log",

LOG_FILES_JSON = r'''
[
    "FedAvg_finetune_100client_label_skew_0.1_5.0_gru.log",
    "FedAvg_lwf_100client_label_skew_0.1_5.0_gru_alpha0.99_temp2.0.log",
    "FedAvg_ewc_100client_label_skew_0.1_5.0_gru_lambda1000.0.log",
    "FedAvg_mas_100client_label_skew_0.1_5.0_gru_lambda200.0.log",
    "FedAvg_icarl_100client_label_skew_0.1_5.0_gru_mem1.0_bce1.log",
    "FedAvg_bic_100client_label_skew_0.1_5.0_gru_mem1.0_alpha0.5_temp2.0_val0.4.log",
    "FedAvg_pass_100client_label_skew_0.1_5.0_gru_lam10.0_g10.0_ps200.log",
    "FedAvg_foster_100client_label_skew_0.1_5.0_gru_mem1.0_b1_0.97_b2_0.97_okd1.0_ce50.log",
    "FedAvg_cbkd_100client_label_skew_0.1_5.0_gru_lam10.0_a10.0_b10.0_ps200.log",
    "FedAvg_feat_100client_label_skew_0.1_5.0_gru_mem200_lam0.1_rho0.9_temp0.5_approx.log",
    "FedAvg_ours3_100client_label_skew_0.1_5.0_gru_mem1.0_kg1.0_gated_eva0.95_blom0.9_synth.log",
    "FedAvg_ours3_100client_label_skew_0.1_5.0_gru_mixed_mem1.0_kg1.0_gated_eva0.95_blom0.3.log",
]
'''
LOG_FILES = []

KW_PATS = [
    (re.compile(r'^lam(?:bda)?(\d+\.?\d*)$'), 'lambda'),
    (re.compile(r'^mem(\d+\.?\d*)$'),          'mem'),
    (re.compile(r'^val(\d+\.?\d*)$'),          'val'),
    (re.compile(r'^temp(\d+\.?\d*)$'),         'temp'),
    (re.compile(r'^bce(\d+\.?\d*)$'),          'bce'),
    (re.compile(r'^g(\d+\.?\d*)$'),            'g'),
    (re.compile(r'^kg(\d+\.?\d*)$'),           'kg'),
    (re.compile(r'^plam(\d+\.?\d*)$'),         'plam'),
    (re.compile(r'^ekd(\d+\.?\d*)$'),          'ekd'),
    (re.compile(r'^ekdl(\d+\.?\d*)$'),         'ekdl'),
    (re.compile(r'^rel(\d+\.?\d*)$'),          'rel'),
    (re.compile(r'^enc(\d+\.?\d*)$'),          'enc'),
    (re.compile(r'^rho(\d+\.?\d*)$'),          'rho'),
    (re.compile(r'^eva(\d+\.?\d*)$'),          'eva'),
    (re.compile(r'^blom(\d+\.?\d*)$'),         'blom'),
    (re.compile(r'^alpha(\d+\.?\d*)$'),        'alpha'),
    (re.compile(r'^meanlogits$'),               'meanlogits'),
]

# Pastel fill colors cycling per strategy group
STRAT_FILLS = [
    'FFD9E1F2', 'FFDDEBF7', 'FFE2EFDA', 'FFFFF2CC',
    'FFFCE4D6', 'FFDEEAF1', 'FFE9D7F5', 'FFFFD7D7',
    'FFE4F5D0', 'FFFDEBD0', 'FFD0EDF5', 'FFFFF0CC',
]
SEP_FILL   = 'FFBDD7EE'
HDR2_FILL  = 'FFD6DCE4'
FINAL_FILL = 'FFFFF2CC'


def _fmtv(s):
    f = float(s)
    return f'{f:g}'


def _make_name(stem, seen):
    tokens = stem.split('_')
    fl, cil = tokens[0], tokens[1]
    base = cil if fl == 'FedAvg' else f'{fl}_{cil}'
    params = []
    for tok in tokens[2:]:
        for pat, kw in KW_PATS:
            m = pat.match(tok)
            if m:
                params.append(kw if m.lastindex is None else f'{kw}={_fmtv(m.group(1))}')
                break
    name = f'{base} ({", ".join(params)})' if params else base
    if name in seen:
        i = 2
        while f'{name} [{i}]' in seen:
            i += 1
        name = f'{name} [{i}]'
    seen.add(name)
    return name


def _metrics(acc, loss=None, f1mac=None, f1mic=None):
    return tuple(None if x is None else float(x) for x in (acc, loss, f1mac, f1mic))


def _num(s):
    return None if str(s).lower() == 'nan' else float(s)


FULL_METRICS = [
    ('acc', 'Acc'),
    ('loss', 'Loss'),
    ('f1mac', 'F1_mac'),
    ('f1mic', 'F1_mic'),
    ('f1w', 'F1_w'),
    ('precmac', 'Prec_mac'),
    ('precmic', 'Prec_mic'),
    ('precw', 'Prec_w'),
    ('recmac', 'Rec_mac'),
    ('recmic', 'Rec_mic'),
    ('recw', 'Rec_w'),
    ('n', 'N'),
]


def _metric_dict(acc=None, loss=None, f1mac=None, f1mic=None, f1w=None,
                 precmac=None, precmic=None, precw=None,
                 recmac=None, recmic=None, recw=None, n=None):
    return {
        'acc': _num(acc) if acc is not None else None,
        'loss': _num(loss) if loss is not None else None,
        'f1mac': _num(f1mac) if f1mac is not None else None,
        'f1mic': _num(f1mic) if f1mic is not None else None,
        'f1w': _num(f1w) if f1w is not None else None,
        'precmac': _num(precmac) if precmac is not None else None,
        'precmic': _num(precmic) if precmic is not None else None,
        'precw': _num(precw) if precw is not None else None,
        'recmac': _num(recmac) if recmac is not None else None,
        'recmic': _num(recmic) if recmic is not None else None,
        'recw': _num(recw) if recw is not None else None,
        'n': int(float(n)) if n is not None else None,
    }


def _line_metrics(line):
    vals = dict(re.findall(r'(Acc|Loss|F1_mac|F1_mic|F1_w|Prec_mac|Prec_mic|Prec_w|Rec_mac|Rec_mic|Rec_w)=((?:nan|[\d.]+))', line))
    precmic = vals.get('Prec_mic') or vals.get('F1_mic') or vals.get('Acc')
    recmic = vals.get('Rec_mic') or vals.get('F1_mic') or vals.get('Acc')
    return _metric_dict(vals.get('Acc'), vals.get('Loss'), vals.get('F1_mac'), vals.get('F1_mic'), vals.get('F1_w'),
                        vals.get('Prec_mac'), precmic, vals.get('Prec_w'),
                        vals.get('Rec_mac'), recmic, vals.get('Rec_w'))


def _clean_log_line(raw, ansi_pat):
    line = ansi_pat.sub('', raw).strip()
    return re.sub(r'^\[\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}\]\s*', '', line)


def _metric_tuple(m):
    return tuple(m.get(k) for k in ('acc', 'f1mac', 'f1mic', 'f1w'))


def _full_metric_tuple(m):
    return tuple(m.get(k) for k, _ in FULL_METRICS)


def _sheet_title(name, seen):
    base = re.sub(r'[\\/*?:\[\]]', '_', name)[:31] or 'Log'
    title = base
    i = 2
    while title in seen:
        suffix = f' {i}'
        title = f'{base[:31 - len(suffix)]}{suffix}'
        i += 1
    seen.add(title)
    return title


def _method_name(name):
    base = name.split(' (', 1)[0]
    return base[:1].upper() + base[1:]


def _parse_log_files_json():
    text = re.sub(r',\s*]', ']', LOG_FILES_JSON)
    return json.loads(text)


def _aggregate_classes(class_metrics, task_map):
    tasks = {}
    class_to_task = {cls: tid for tid, classes in task_map.items() for cls in classes}
    for cls, m in class_metrics.items():
        task = class_to_task.get(cls)
        if task is None:
            continue
        tasks.setdefault(task, []).append(m)
    out = {}
    for task, rows in tasks.items():
        support = sum(m['n'] for m in rows if m.get('n') is not None)
        acc = sum(m['acc'] * m['n'] for m in rows if m.get('acc') is not None and m.get('n') is not None) / support
        f1mac = sum(m['f1'] for m in rows if m.get('f1') is not None) / len(rows)
        f1mic = acc
        f1w = sum(m['f1'] * m['n'] for m in rows if m.get('f1') is not None and m.get('n') is not None) / support
        precmac = sum(m['prec'] for m in rows if m.get('prec') is not None) / len(rows)
        precmic = acc
        precw = sum(m['prec'] * m['n'] for m in rows if m.get('prec') is not None and m.get('n') is not None) / support
        recmac = sum(m['rec'] for m in rows if m.get('rec') is not None) / len(rows)
        recmic = acc
        recw = sum(m['rec'] * m['n'] for m in rows if m.get('rec') is not None and m.get('n') is not None) / support
        out[task] = _metric_dict(acc, None, f1mac, f1mic, f1w,
                     precmac, precmic, precw,
                     recmac, recmic, recw,
                     n=support)
    return out


def _parse_log(path, task_map):
    tasks = {}
    merged = {}
    final = None
    per_classes = {}
    ansi_pat = re.compile(r'\x1b\[[0-9;]*m')
    num = r'(?:nan|[\d.]+)'
    old_task_eval_pat = re.compile(r'^EVAL \| T(\d+) \| R(\d+) \| Task(\d+) \| ')
    eval_pat = re.compile(r'^EVAL \| T(\d+) \| R(\d+) \| Acc=')
    final_pat = re.compile(r'^FINAL \| all \d+ classes \|')
    per_task_pat = re.compile(r'^EVAL \| T(\d+) \| R(\d+) \| per-task \| (.+)$')
    task_item_pat = re.compile(rf'T(\d+):acc=({num}),f1=({num}),n=({num})')
    per_class_pat = re.compile(r'^EVAL \| T(\d+) \| R(\d+) \| per-class \| (.+)$')
    class_item_pat = re.compile(rf'C(\d+):acc=({num}),prec=({num}),rec=({num}),f1=({num}),n=({num})')
    for raw in path.read_text(encoding='utf-8', errors='replace').splitlines():
        line = _clean_log_line(raw, ansi_pat)
        m = old_task_eval_pat.match(line)
        if m:
            train_task = int(m.group(1))
            R = int(m.group(2))
            eval_task = int(m.group(3))
            tasks.setdefault(train_task, {}).setdefault(R, {})[eval_task] = _line_metrics(line)
            continue
        m = eval_pat.match(line)
        if m:
            train_task = int(m.group(1))
            R = int(m.group(2))
            merged.setdefault(train_task, {})[R] = _line_metrics(line)
            tasks.setdefault(train_task, {})
            continue
        m = per_class_pat.match(line)
        if m:
            train_task = int(m.group(1))
            R = int(m.group(2))
            class_metrics = {}
            for item in class_item_pat.finditer(m.group(3)):
                class_metrics[int(item.group(1))] = {
                    'acc': _num(item.group(2)),
                    'prec': _num(item.group(3)),
                    'rec': _num(item.group(4)),
                    'f1': _num(item.group(5)),
                    'n': int(float(item.group(6))),
                }
            per_classes[(train_task, R)] = class_metrics
            for eval_task, vals in _aggregate_classes(class_metrics, task_map).items():
                tasks.setdefault(train_task, {}).setdefault(R, {})[eval_task] = vals
            continue
        m = per_task_pat.match(line)
        if m:
            train_task = int(m.group(1))
            R = int(m.group(2))
            class_tasks = _aggregate_classes(per_classes.get((train_task, R), {}), task_map)
            for item in task_item_pat.finditer(m.group(3)):
                eval_task = int(item.group(1))
                support = item.group(4)
                class_task = class_tasks.get(eval_task, {})
                f1w = class_task.get('f1w')
                tasks.setdefault(train_task, {}).setdefault(R, {})[eval_task] = _metric_dict(
                    item.group(2), None, item.group(3), item.group(2), f1w,
                    class_task.get('precmac'), class_task.get('precmic'), class_task.get('precw'),
                    class_task.get('recmac'), class_task.get('recmic'), class_task.get('recw'),
                    n=support)
            continue
        m = final_pat.match(line)
        if m:
            final = _line_metrics(line)
    if final is None and merged:
        train_task = max(merged)
        R = max(merged[train_task])
        final = merged[train_task][R]
    return {'tasks': tasks, 'merged': merged, 'final': final}


def _load_task_map():
    task_map = {}
    for line in TASK_FILE.read_text().splitlines():
        line = line.strip()
        if not line or line.startswith('#'):
            continue
        m = re.match(r'^(\d+):\s*\[([\d,\s]+)\]', line)
        if m:
            task_map[int(m.group(1))] = [int(x.strip()) for x in m.group(2).split(',') if x.strip()]
    return task_map


def _load_label_map():
    lmap = {}
    for line in LABEL_FILE.read_text().splitlines():
        if ':' in line:
            idx, name = line.split(':', 1)
            lmap[int(idx.strip())] = name.strip()
    return lmap


def _cell_fill(ws, row, col, value=None, fill_hex=None, bold=False, center=True, wrap=False, num_fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    if fill_hex:
        c.fill = PatternFill('solid', fgColor=fill_hex)
    c.font = Font(bold=bold)
    align_kw = {'vertical': 'center'}
    if center:
        align_kw['horizontal'] = 'center'
    if wrap:
        align_kw['wrap_text'] = True
    c.alignment = Alignment(**align_kw)
    if num_fmt:
        c.number_format = num_fmt
    return c


def _sep_row(ws, row, n_cols, text, fill=SEP_FILL):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
    c = ws.cell(row=row, column=1, value=text)
    c.fill = PatternFill('solid', fgColor=fill)
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal='left', vertical='center')


def build_comparison(ws, all_data, names, task_map):
    n = len(names)
    METRICS = ['Acc', 'F1_mac', 'F1_mic', 'F1_w']
    NM = len(METRICS)
    total_cols = 1 + NM * n

    ws.row_dimensions[1].height = 45
    _cell_fill(ws, 1, 1, 'Round', fill_hex=HDR2_FILL, bold=True)
    for i, name in enumerate(names):
        col = 2 + NM * i
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + NM - 1)
        _cell_fill(ws, 1, col, name, fill_hex=STRAT_FILLS[i % len(STRAT_FILLS)], bold=True, wrap=True)

    ws.row_dimensions[2].height = 20
    _cell_fill(ws, 2, 1, 'Task/Round', fill_hex=HDR2_FILL, bold=True, center=True)
    for i in range(n):
        fill = STRAT_FILLS[i % len(STRAT_FILLS)]
        for mi, m in enumerate(METRICS):
            _cell_fill(ws, 2, 2 + NM * i + mi, m, fill_hex=fill, bold=True)

    all_train_tasks = sorted({t for d in all_data.values() for t in set(d['tasks']) | set(d.get('merged', {}))})
    cur = 3

    for train_task in all_train_tasks:
        classes = task_map.get(train_task, [])
        _sep_row(ws, cur, total_cols, f"Training Task {train_task}  ·  classes {classes}")
        cur += 1

        all_rounds = sorted({R
                             for d in all_data.values()
                             for R in set(d['tasks'].get(train_task, {})) | set(d.get('merged', {}).get(train_task, {}))})
        for R in all_rounds:
            ws.cell(row=cur, column=1, value=f'T{train_task} R{R}').alignment = Alignment(horizontal='center')
            for i, nm in enumerate(names):
                entry = all_data[nm].get('merged', {}).get(train_task, {}).get(R)
                if entry:
                    for mi in range(NM):
                        val = _metric_tuple(entry)[mi]
                        if val is not None:
                            _cell_fill(ws, cur, 2 + NM * i + mi, val, num_fmt='0.0000')
            cur += 1


    # Final row
    _sep_row(ws, cur, total_cols, "Final (all classes)", fill=FINAL_FILL)
    cur += 1
    ws.cell(row=cur, column=1, value='Final').alignment = Alignment(horizontal='center')
    for i, nm in enumerate(names):
        f = all_data[nm].get('final')
        if f:
            for mi in range(NM):
                val = _metric_tuple(f)[mi]
                if val is not None:
                    _cell_fill(ws, cur, 2 + NM * i + mi, val, num_fmt='0.0000')

    ws.column_dimensions['A'].width = 18
    for i in range(n):
        for mi in range(NM):
            ws.column_dimensions[get_column_letter(2 + NM * i + mi)].width = 9

    ws.freeze_panes = 'B3'


def build_log_sheet(ws, data, name, task_map):
    total_cols = 2 + len(FULL_METRICS)
    title = _method_name(name)
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    _cell_fill(ws, 1, 1, title, fill_hex=STRAT_FILLS[0], bold=True, wrap=True)

    headers = ['Training', 'Round'] + [label for _, label in FULL_METRICS]
    for col, header in enumerate(headers, 1):
        _cell_fill(ws, 2, col, header, fill_hex=HDR2_FILL, bold=True, wrap=True)

    all_train_tasks = sorted(data.get('merged', {}))
    cur = 3
    for train_task in all_train_tasks:
        classes = task_map.get(train_task, [])
        _sep_row(ws, cur, total_cols, f'Training Task {train_task}  ·  classes {classes}')
        cur += 1

        for R in sorted(data.get('merged', {}).get(train_task, {})):
            cur = _write_detail_row(ws, cur, train_task, R, data['merged'][train_task][R])

    final = data.get('final')
    if final:
        _sep_row(ws, cur, total_cols, 'Final (all classes)', fill=FINAL_FILL)
        cur += 1
        cur = _write_detail_row(ws, cur, None, None, final)

    widths = [12, 9] + [10] * len(FULL_METRICS)
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = 'C3'


def _write_detail_row(ws, row, train_task, R, entry):
    ws.cell(row=row, column=1, value=None if train_task is None else f'T{train_task}').alignment = Alignment(horizontal='center')
    ws.cell(row=row, column=2, value=None if R is None else f'R{R}').alignment = Alignment(horizontal='center')
    for mi, val in enumerate(_full_metric_tuple(entry), 3):
        if val is not None:
            fmt = '0' if FULL_METRICS[mi - 3][0] == 'n' else '0.0000'
            _cell_fill(ws, row, mi, val, num_fmt=fmt)
    return row + 1


def _client_id_from_label(s):
    m = re.search(r'client_(\d+)', str(s))
    return int(m.group(1)) if m else None


def _task_active_n(n_clients, num_tasks, task_id):
    if num_tasks != 7:
        return n_clients
    return max(1, int(n_clients * (0.4 + task_id * 0.1)))


def build_task_partition(ws, task_map, label_map):
    dist_wb = openpyxl.load_workbook(DIST_XLSX)
    dist_ws = dist_wb.active
    rows = list(dist_ws.iter_rows(values_only=True))
    header    = rows[0]
    data_rows = rows[1:]

    # Map class id → column index in the original header
    class_col = {}
    extra_cols = []   # (col_idx, header_value) for non-class, non-client columns (e.g. Total)
    for i, h in enumerate(header):
        if i == 0:
            continue   # Client column
        try:
            class_col[int(h)] = i
        except (ValueError, TypeError):
            if h is not None:
                extra_cols.append((i, h))

    # Build ordered class list from task map
    ordered_classes = []
    task_col_ranges = []   # (tid, start_spreadsheet_col, end_spreadsheet_col)
    next_col = 2           # spreadsheet col; col 1 = Client
    for tid in sorted(task_map.keys()):
        start_col = next_col
        for cls in task_map[tid]:
            ordered_classes.append(cls)
            next_col += 1
        task_col_ranges.append((tid, start_col, next_col - 1))

    # Extra columns (Total etc.) follow class columns
    extra_start = next_col

    # Build class → task id lookup
    cls_to_tid = {cls: tid for tid, cls_list in task_map.items() for cls in cls_list}
    num_tasks = len(task_map)
    n_clients = len([drow for drow in data_rows if _client_id_from_label(drow[0]) is not None])

    # Row 1: task group merged headers
    ws.row_dimensions[1].height = 22
    _cell_fill(ws, 1, 1, 'Client', bold=True)
    for tid, sc, ec in task_col_ranges:
        if sc < ec:
            ws.merge_cells(start_row=1, start_column=sc, end_row=1, end_column=ec)
        fill = STRAT_FILLS[tid % len(STRAT_FILLS)]
        c = ws.cell(row=1, column=sc, value=f'Task {tid}')
        c.fill = PatternFill('solid', fgColor=fill)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal='center', vertical='center')
    for i, (_, eh) in enumerate(extra_cols):
        ws.cell(row=1, column=extra_start + i, value=str(eh)).font = Font(bold=True)

    # Row 2: class index + name
    ws.row_dimensions[2].height = 50
    _cell_fill(ws, 2, 1, 'Client', fill_hex=HDR2_FILL, bold=True)
    for i, cls_id in enumerate(ordered_classes):
        col  = i + 2
        tid  = cls_to_tid.get(cls_id, 0)
        fill = STRAT_FILLS[tid % len(STRAT_FILLS)]
        lbl  = f"{cls_id}\n{label_map.get(cls_id, str(cls_id))}"
        c = ws.cell(row=2, column=col, value=lbl)
        c.fill = PatternFill('solid', fgColor=fill)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = 13
    for i, (_, eh) in enumerate(extra_cols):
        col = extra_start + i
        c = ws.cell(row=2, column=col, value=str(eh))
        c.fill = PatternFill('solid', fgColor=HDR2_FILL)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(col)].width = 14

    # Data rows
    for r_idx, drow in enumerate(data_rows):
        row = r_idx + 3
        ws.cell(row=row, column=1, value=drow[0])
        cid = _client_id_from_label(drow[0])
        for i, cls_id in enumerate(ordered_classes):
            orig_idx = class_col.get(cls_id)
            val = drow[orig_idx] if orig_idx is not None else None
            tid = cls_to_tid.get(cls_id)
            if cid is not None and tid is not None and cid >= _task_active_n(n_clients, num_tasks, tid):
                val = None
            c = ws.cell(row=row, column=i + 2, value=val)
            c.alignment = Alignment(horizontal='right')
        for i, (orig_idx, _) in enumerate(extra_cols):
            c = ws.cell(row=row, column=extra_start + i, value=drow[orig_idx])
            c.alignment = Alignment(horizontal='right')

    ws.column_dimensions['A'].width = 16
    ws.freeze_panes = 'B3'


def main():
    task_map = _load_task_map()
    label_map = _load_label_map()

    if len(sys.argv) > 1:
        files = [RESULTS_DIR / f for f in sys.argv[1:]]
    elif _parse_log_files_json():
        LOG_FILES[:] = _parse_log_files_json()
        files = [RESULTS_DIR / f for f in LOG_FILES]
    else:
        files = sorted(RESULTS_DIR.glob('*.log'))

    missing = [f for f in files if not Path(f).exists()]
    if missing:
        print('Missing logs:')
        for f in missing:
            print(f'  {f}')
        sys.exit(1)

    seen  = set()
    names = []
    all_data = {}
    for f in files:
        nm = _make_name(Path(f).stem, seen)
        names.append(nm)
        all_data[nm] = _parse_log(Path(f), task_map)

    wb       = openpyxl.Workbook()
    ws_comp  = wb.active
    ws_comp.title = 'Comparison'
    build_comparison(ws_comp, all_data, names, task_map)

    sheet_names = {ws_comp.title}
    for nm in names:
        ws_log = wb.create_sheet(_sheet_title(_method_name(nm), sheet_names))
        build_log_sheet(ws_log, all_data[nm], nm, task_map)

    ws_part = wb.create_sheet('Task Partition')
    build_task_partition(ws_part, task_map, label_map)

    out = RESULTS_DIR / 'cil_comparison.xlsx'
    wb.save(out)
    print(f'Saved: {out}')


if __name__ == '__main__':
    main()
